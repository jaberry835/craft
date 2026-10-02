import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from markdown_it import MarkdownIt
from openpyxl import Workbook, load_workbook
from openpyxl.cell import Cell
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
from openpyxl.workbook import Workbook as WorkbookObject

from officeautomation.errors import ServiceError
from officeautomation.files.local import LocalFileStore
from officeautomation.models.docx import ApplyResult, FileOperationResult, MarkdownResult
from officeautomation.models.xlsx import MarkdownToXlsxRequest, XlsxApplyRequest

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PLACEHOLDER_PATTERN = re.compile(r"\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|«[^»]+»|<<[^<>]+>>")


class XlsxService:
    def __init__(self, store: LocalFileStore) -> None:
        self.store = store
        self.markdown = MarkdownIt("commonmark").enable("table")

    async def from_markdown(
        self, owner: str, request: MarkdownToXlsxRequest
    ) -> FileOperationResult:
        markdown = self._load_markdown(owner, request.markdown, request.markdown_file_id)
        tables = self._extract_tables(markdown)
        if not tables:
            raise ServiceError("INVALID_DOCUMENT", "Markdown contains no tables", 422)
        if request.template_file_id:
            _, template_path = self.store.open_content(owner, request.template_file_id)
            workbook = self._load_workbook(template_path)
        else:
            workbook = Workbook()
            workbook.remove(workbook.active)
        used_names: set[str] = set(workbook.sheetnames)
        for table_index, (suggested_name, rows) in enumerate(tables, start=1):
            name = self._unique_sheet_name(suggested_name or f"Sheet{table_index}", used_names)
            if request.template_file_id and name in workbook.sheetnames:
                worksheet = workbook[name]
            else:
                worksheet = workbook.create_sheet(name)
            for row_index, values in enumerate(rows, start=1):
                for column_index, value in enumerate(values, start=1):
                    worksheet.cell(row_index, column_index, value=value)
        self._set_recalculation(workbook)
        return await self._save_workbook(
            owner,
            workbook,
            request.output_name,
            request.template_file_id or request.markdown_file_id,
        )

    async def to_markdown(
        self,
        owner: str,
        file_id: str,
        sheets: list[str] | None,
        values_or_formulas: str,
        output_mode: str,
    ) -> MarkdownResult:
        stored, path = self.store.open_content(owner, file_id)
        workbook = self._load_workbook(
            path,
            data_only=values_or_formulas == "values",
            read_only=True,
        )
        selected = sheets or workbook.sheetnames
        missing = [name for name in selected if name not in workbook.sheetnames]
        if missing:
            raise ServiceError(
                "ANCHOR_NOT_FOUND", "One or more sheets were not found", 404, {"sheets": missing}
            )
        sections: list[str] = []
        for name in selected:
            worksheet = workbook[name]
            rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            while rows and not any(value is not None for value in rows[-1]):
                rows.pop()
            if not rows:
                continue
            width = max(
                (index + 1 for row in rows for index, value in enumerate(row) if value is not None),
                default=1,
            )
            normalized = [[self._format_value(value) for value in row[:width]] for row in rows]
            sections.extend([f"## {name}", "", self._rows_to_markdown(normalized), ""])
        markdown = "\n".join(sections).rstrip() + "\n"
        if output_mode == "file" or (output_mode == "auto" and len(markdown.encode()) > 100_000):
            reference = await self.store.save_bytes(
                owner,
                f"{Path(stored.name).stem}.md",
                "text/markdown",
                markdown.encode(),
                parent_file_id=file_id,
            )
            return MarkdownResult(markdown_file=reference)
        return MarkdownResult(markdown=markdown)

    def inspect(self, owner: str, file_id: str, sample_rows: int) -> dict[str, Any]:
        _, path = self.store.open_content(owner, file_id)
        workbook = self._load_workbook(path, data_only=False)
        sheets: list[dict[str, Any]] = []
        placeholders: list[dict[str, str]] = []
        formula_count = 0
        at_risk: list[str] = []
        for worksheet in workbook.worksheets:
            samples: list[list[Any]] = []
            for row in worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, sample_rows)):
                samples.append([cell.value for cell in row])
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
                    if isinstance(cell.value, str):
                        for token in PLACEHOLDER_PATTERN.findall(cell.value):
                            placeholders.append(
                                {"token": token, "cell": f"{worksheet.title}!{cell.coordinate}"}
                            )
            if worksheet._charts:
                at_risk.append(f"{worksheet.title}:charts")
            if worksheet._images:
                at_risk.append(f"{worksheet.title}:images")
            sheets.append(
                {
                    "name": worksheet.title,
                    "visibility": worksheet.sheet_state,
                    "used_range": worksheet.calculate_dimension(),
                    "merged_cells": [str(item) for item in worksheet.merged_cells.ranges],
                    "freeze_panes": str(worksheet.freeze_panes) if worksheet.freeze_panes else None,
                    "tables": [
                        {"name": table.name, "range": table.ref}
                        for table in worksheet.tables.values()
                    ],
                    "sample_rows": samples,
                }
            )
        named_ranges = [name for name in workbook.defined_names]
        return {
            "file_id": file_id,
            "sheets": sheets,
            "named_ranges": named_ranges,
            "placeholders": placeholders,
            "formula_count": formula_count,
            "features_at_risk": at_risk,
        }

    async def apply(self, owner: str, request: XlsxApplyRequest) -> ApplyResult:
        _, path = self.store.open_content(owner, request.file_id)
        workbook = self._load_workbook(path, data_only=False)
        risks = self._feature_risks(workbook)
        if risks and not request.allow_feature_loss:
            raise ServiceError(
                "FEATURE_LOSS_BLOCKED",
                "Workbook contains features that may be changed by openpyxl",
                409,
                {"features": risks},
            )
        results: list[dict[str, Any]] = []
        for operation in request.operations:
            count = 0
            if operation.type == "set_cells":
                if not isinstance(operation.values, dict):
                    raise ServiceError(
                        "INVALID_DOCUMENT", "set_cells requires an address-to-value map", 422
                    )
                for address, value in operation.values.items():
                    worksheet, coordinate = self._resolve_address(workbook, address)
                    self._set_cell(worksheet[coordinate], value, request.allow_formulas)
                    count += 1
            elif operation.type == "write_range":
                if (
                    not isinstance(operation.values, list)
                    or not operation.sheet
                    or not operation.start_cell
                ):
                    raise ServiceError(
                        "INVALID_DOCUMENT",
                        "write_range requires sheet, start_cell, and values",
                        422,
                    )
                worksheet = self._sheet(workbook, operation.sheet)
                column_letters, start_row = coordinate_from_string(operation.start_cell)
                start_column = column_index_from_string(column_letters)
                for row_offset, values in enumerate(operation.values):
                    for column_offset, value in enumerate(values):
                        self._set_cell(
                            worksheet.cell(start_row + row_offset, start_column + column_offset),
                            value,
                            request.allow_formulas,
                        )
                        count += 1
            elif operation.type == "replace_placeholders":
                replacements = operation.replacements or {}
                for worksheet in workbook.worksheets:
                    for row in worksheet.iter_rows():
                        for cell in row:
                            if not isinstance(cell.value, str):
                                continue
                            updated = cell.value
                            for token, value in replacements.items():
                                matches = updated.count(token)
                                if matches:
                                    updated = updated.replace(token, str(value))
                                    count += matches
                            if updated != cell.value:
                                self._set_cell(cell, updated, False)
                if count == 0:
                    raise ServiceError("ANCHOR_NOT_FOUND", "No placeholders were found", 404)
            elif operation.type == "add_sheet":
                if not operation.name:
                    raise ServiceError("INVALID_DOCUMENT", "add_sheet requires name", 422)
                workbook.create_sheet(operation.name)
                count = 1
            elif operation.type == "rename_sheet":
                if not operation.name or not operation.new_name:
                    raise ServiceError(
                        "INVALID_DOCUMENT", "rename_sheet requires name and new_name", 422
                    )
                self._sheet(workbook, operation.name).title = operation.new_name
                count = 1
            elif operation.type == "delete_sheet":
                if not operation.name:
                    raise ServiceError("INVALID_DOCUMENT", "delete_sheet requires name", 422)
                if len(workbook.sheetnames) == 1:
                    raise ServiceError(
                        "INVALID_DOCUMENT", "A workbook must contain at least one sheet", 422
                    )
                workbook.remove(self._sheet(workbook, operation.name))
                count = 1
            results.append({"type": operation.type, "status": "applied", "matches": count})
        self._set_recalculation(workbook)
        saved = await self._save_workbook(owner, workbook, request.output_name, request.file_id)
        return ApplyResult(file=saved.file, warnings=saved.warnings, operations=results)

    def _load_markdown(self, owner: str, inline: str | None, file_id: str | None) -> str:
        if inline is not None:
            return inline
        if not file_id:
            raise ServiceError("INVALID_DOCUMENT", "Markdown input is required", 422)
        _, path = self.store.open_content(owner, file_id)
        return path.read_text(encoding="utf-8")

    def _extract_tables(self, markdown: str) -> list[tuple[str | None, list[list[str]]]]:
        tokens = self.markdown.parse(markdown)
        heading: str | None = None
        tables: list[tuple[str | None, list[list[str]]]] = []
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "heading_open":
                heading = tokens[index + 1].content
                index += 3
                continue
            if token.type == "table_open":
                rows: list[list[str]] = []
                row: list[str] | None = None
                index += 1
                while tokens[index].type != "table_close":
                    current = tokens[index]
                    if current.type == "tr_open":
                        row = []
                    elif current.type in {"th_open", "td_open"} and row is not None:
                        row.append(tokens[index + 1].content)
                    elif current.type == "tr_close" and row is not None:
                        rows.append(row)
                        row = None
                    index += 1
                tables.append((heading, rows))
            index += 1
        return tables

    @staticmethod
    def _load_workbook(path: Path, **options: Any) -> WorkbookObject:
        return load_workbook(BytesIO(path.read_bytes()), **options)

    async def _save_workbook(
        self,
        owner: str,
        workbook: WorkbookObject,
        name: str,
        parent_file_id: str | None,
    ) -> FileOperationResult:
        output = BytesIO()
        workbook.save(output)
        output_name = name if name.lower().endswith(".xlsx") else f"{name}.xlsx"
        reference = await self.store.save_bytes(
            owner,
            output_name,
            XLSX_CONTENT_TYPE,
            output.getvalue(),
            parent_file_id=parent_file_id,
        )
        return FileOperationResult(file=reference)

    @staticmethod
    def _set_cell(cell: Cell, value: Any, allow_formulas: bool) -> None:
        cell.value = value
        if isinstance(value, str) and value.startswith("=") and not allow_formulas:
            cell.data_type = "s"

    @staticmethod
    def _resolve_address(workbook: WorkbookObject, address: str) -> tuple[Any, str]:
        if "!" not in address:
            raise ServiceError(
                "INVALID_DOCUMENT", f"Cell address must include a sheet: {address}", 422
            )
        sheet_name, coordinate = address.rsplit("!", 1)
        return XlsxService._sheet(workbook, sheet_name.strip("'")), coordinate

    @staticmethod
    def _sheet(workbook: WorkbookObject, name: str) -> Any:
        if name not in workbook.sheetnames:
            raise ServiceError("ANCHOR_NOT_FOUND", f"Sheet was not found: {name}", 404)
        return workbook[name]

    @staticmethod
    def _feature_risks(workbook: WorkbookObject) -> list[str]:
        risks: list[str] = []
        for worksheet in workbook.worksheets:
            if worksheet._charts:
                risks.append(f"{worksheet.title}:charts")
            if worksheet._images:
                risks.append(f"{worksheet.title}:images")
        return risks

    @staticmethod
    def _set_recalculation(workbook: WorkbookObject) -> None:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True

    @staticmethod
    def _format_value(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        return str(value).replace("|", "\\|").replace("\n", "<br>")

    @staticmethod
    def _rows_to_markdown(rows: list[list[str]]) -> str:
        width = max(len(row) for row in rows)
        normalized = [row + [""] * (width - len(row)) for row in rows]
        output = [
            "| " + " | ".join(normalized[0]) + " |",
            "| " + " | ".join(["---"] * width) + " |",
        ]
        output.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
        return "\n".join(output)

    @staticmethod
    def _unique_sheet_name(candidate: str, used: set[str]) -> str:
        cleaned = re.sub(r"[\\/*?:\[\]]", "_", candidate).strip()[:31] or "Sheet"
        name = cleaned
        suffix = 2
        while name in used:
            marker = f" {suffix}"
            name = f"{cleaned[: 31 - len(marker)]}{marker}"
            suffix += 1
        used.add(name)
        return name
